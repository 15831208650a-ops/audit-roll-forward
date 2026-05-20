#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
详细检查C货币资金的Lead Sheet结构
"""

import openpyxl
import os

def inspect_C_lead(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet_name = "C.00 Lead"
    if sheet_name not in wb.sheetnames:
        print(f"找不到Sheet: {sheet_name}")
        return

    ws = wb[sheet_name]
    print(f"Sheet: {sheet_name}, 尺寸: {ws.max_row}x{ws.max_column}\n")

    # 查找"期末审定数"所在行
    header_row = None
    for row in range(1, min(80, ws.max_row + 1)):
        for col in range(1, min(20, ws.max_column + 1)):
            val = ws.cell(row=row, column=col).value
            if val and "期末审定数" in str(val):
                header_row = row
                print(f"找到'期末审定数'在第{row}行, 第{col}列")
                # 显示该行的前15列
                row_data = []
                for c in range(1, min(16, ws.max_column + 1)):
                    v = ws.cell(row=row, column=c).value
                    if v:
                        row_data.append(f"C{c}={str(v)[:30]}")
                print(f"  表头行: {', '.join(row_data)}")
                break
        if header_row:
            break

    # 查找"For Disclosure"或"披露"相关行
    print("\n查找For Disclosure或披露相关行:")
    for row in range(1, ws.max_row + 1):
        for col in range(1, 20):
            val = ws.cell(row=row, column=col).value
            if val and any(x in str(val) for x in ["Disclosure", "披露", "表2", "For"]):
                print(f"  R{row}C{col}: {str(val)[:80]}")

    # 查找PMTE/CRA相关行
    print("\n查找PMTE/CRA相关行:")
    for row in range(1, 20):
        for col in range(1, 10):
            val = ws.cell(row=row, column=col).value
            if val and any(x in str(val) for x in ["PMTE", "CRA", "TE", "SAD", "RP", "层级"]):
                print(f"  R{row}C{col}: {str(val)[:80]}")

    # 查找日期相关
    print("\n查找日期相关行:")
    for row in range(1, 40):
        for col in range(1, 10):
            val = ws.cell(row=row, column=col).value
            if val and any(x in str(val) for x in ["202", "日期", "期末", "资产负债表"]):
                print(f"  R{row}C{col}: {str(val)[:80]}")

    wb.close()

if __name__ == "__main__":
    inspect_C_lead(r"F:\AI专用\AI工作区\底稿roll forward\output_final\C SWP 货币资金 20261231 东风鸿远工程咨询有限公司.xlsx")
