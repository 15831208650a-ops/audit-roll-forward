import openpyxl

pmte_path = "输入；参考资料.xlsx"
wb = openpyxl.load_workbook(pmte_path, data_only=True)
ws = wb["CRA"]

print("=== CRA 表完整结构 ===")
print(f"总行数: {ws.max_row}, 总列数: {ws.max_column}")

# 打印前50行的关键列
for row in range(1, 51):
    for col in range(1, 12):
        val = ws.cell(row=row, column=col).value
        if val:
            print(f"  ({row},{col}): {str(val)[:80]}")

wb.close()
