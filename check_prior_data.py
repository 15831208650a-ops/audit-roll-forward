#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查上年底稿的数据情况（data_only=True vs data_only=False）
"""

import openpyxl

def check_prior_C(prior_path):
    print(f"检查上年底稿: {prior_path}\n")

    # data_only=True（读取计算值）
    wb_values = openpyxl.load_workbook(prior_path, data_only=True)
    ws_values = wb_values["C.00 Lead"]

    # data_only=False（读取公式）
    wb_formula = openpyxl.load_workbook(prior_path, data_only=False)
    ws_formula = wb_formula["C.00 Lead"]

    # 查找表头行
    header_row = None
    for row in range(1, 80):
        for col in range(1, 20):
            val = ws_values.cell(row=row, column=col).value
            if val and "期末审定数" in str(val):
                header_row = row
                print(f"找到'期末审定数'在第{row}行, 第{col}列")
                break
        if header_row:
            break

    if not header_row:
        print("找不到表头行")
        return

    # 检查第38行（现金行）
    print(f"\n=== 第{header_row+1}行（现金行）数据对比 ===")
    for col in range(9, 15):
        val_true = ws_values.cell(row=header_row+1, column=col).value
        val_false = ws_formula.cell(row=header_row+1, column=col).value
        print(f"  C{col}: data_only=True={val_true}, data_only=False={val_false}")

    # 检查For Disclosure区域
    print("\n=== For Disclosure区域 ===")
    for row in range(60, 76):
        val_c2 = ws_values.cell(row=row, column=2).value
        if val_c2:
            print(f"  R{row}C2: {val_c2}")
            # 显示该行的C3-C6
            for col in range(3, 7):
                val_true = ws_values.cell(row=row, column=col).value
                val_false = ws_formula.cell(row=row, column=col).value
                if val_true is not None or val_false is not None:
                    print(f"    C{col}: data_only=True={val_true}, data_only=False={val_false}")

    wb_values.close()
    wb_formula.close()

if __name__ == "__main__":
    check_prior_C(r"F:\AI专用\AI工作区\底稿roll forward\上年底稿文件夹\C SWP 货币资金 20251231 东风鸿远工程咨询有限公司.xlsx")
