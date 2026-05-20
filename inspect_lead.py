#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查所有科目的Lead Sheet结构
"""

import os
import openpyxl

def inspect_lead_sheet(filepath):
    """检查Lead Sheet的关键位置"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    result = {"filename": os.path.basename(filepath)}

    # 找到Lead Sheet
    lead_sheet = None
    for sheet_name in wb.sheetnames:
        if "Lead" in sheet_name:
            lead_sheet = sheet_name
            break

    if not lead_sheet:
        result["error"] = "找不到Lead Sheet"
        return result

    ws = wb[lead_sheet]
    result["lead_sheet"] = lead_sheet
    result["dimensions"] = f"{ws.max_row}x{ws.max_column}"

    # 查找"期末审定数"和"期初数"的位置
    for row in range(1, min(80, ws.max_row + 1)):
        for col in range(1, min(20, ws.max_column + 1)):
            val = ws.cell(row=row, column=col).value
            if val:
                text = str(val)
                if "期末审定数" in text:
                    result["期末审定数"] = {"row": row, "col": col}
                if any(x in text for x in ["期初数", "年初余额", "期初余额"]):
                    result["期初数"] = {"row": row, "col": col}
                if "客户名称" in text:
                    result["客户名称"] = {"row": row, "col": col}
                if "期末" in text and "202" in text:
                    result["资产负债表日"] = {"row": row, "col": col, "text": text}
                if "分析日期" in text:
                    result["分析日期"] = {"row": row, "col": col}
                if "可容忍误差" in text or "TE" in text:
                    result["TE"] = {"row": row, "col": col}
                if "名义金额" in text or "SAD" in text:
                    result["SAD"] = {"row": row, "col": col}

    # 读取Lead Sheet表头前几行
    header_preview = {}
    for row in range(1, min(12, ws.max_row + 1)):
        row_data = []
        for col in range(1, min(8, ws.max_column + 1)):
            val = ws.cell(row=row, column=col).value
            if val:
                text = str(val)[:50]
                if any(x in text for x in ["客户名称", "期末", "分析日期", "可容忍", "名义金额", "适用会计", "记账本位"]):
                    row_data.append(f"C{col}={text}")
        if row_data:
            header_preview[f"R{row}"] = row_data
    result["header_preview"] = header_preview

    wb.close()
    return result


def main():
    base_dir = r"F:\AI专用\AI工作区\底稿roll forward"
    files = [f for f in os.listdir(base_dir) if f.endswith('.xlsx') and not f.startswith('~') and not f.startswith('输入')]

    for f in sorted(files):
        filepath = os.path.join(base_dir, f)
        try:
            result = inspect_lead_sheet(filepath)
            print(f"\n{'='*80}")
            print(f"文件: {result['filename']}")
            if "error" in result:
                print(f"  错误: {result['error']}")
                continue

            print(f"  Lead Sheet: {result['lead_sheet']}")
            print(f"  尺寸: {result['dimensions']}")

            for key in ["客户名称", "资产负债表日", "分析日期", "TE", "SAD", "期末审定数", "期初数"]:
                if key in result:
                    info = result[key]
                    print(f"  [{key}] 行{info['row']}, 列{info['col']}", end="")
                    if "text" in info:
                        print(f" = {info['text']}")
                    else:
                        print()

        except Exception as e:
            print(f"\n文件 {f} 检查失败: {e}")


if __name__ == "__main__":
    main()
