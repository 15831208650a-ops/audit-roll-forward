#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查模板中C5的值和D15:D19的公式
"""

import openpyxl

def check_template(template_path):
    wb = openpyxl.load_workbook(template_path, data_only=False)
    ws = wb["C.00 Lead"]

    print(f"=== C.00 Lead Sheet ===")
    print(f"C5的值: {ws.cell(row=5, column=3).value}")
    print(f"C5的数据类型: {type(ws.cell(row=5, column=3).value)}")

    print("\n=== D15:D19 的公式 ===")
    for row in range(15, 20):
        cell = ws.cell(row=row, column=4)
        c_val = ws.cell(row=row, column=3).value
        print(f"R{row}: C={c_val}, D={cell.value}")

    # 也检查C15:C19的值（如果它们是手动输入的）
    print("\n=== C15:C19 的值 ===")
    for row in range(15, 20):
        val = ws.cell(row=row, column=3).value
        print(f"R{row}C3: {val}")

    wb.close()

check_template(r"F:\AI专用\AI工作区\底稿roll forward\templates\C SWP 货币资金 202YMMDD XYZ公司.xlsx")
