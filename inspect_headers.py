import openpyxl
import json

# 检查上年底稿的表头
prior_path = "输入：K1 固定资产 20251231六六六有限公司.xlsx"
wb = openpyxl.load_workbook(prior_path, data_only=False)

result = {}

# K.00 Lead Sheet
ws = wb["K.00 Lead Sheet"]
k00_headers = []
for row in range(1, 15):
    for col in range(1, 15):
        val = ws.cell(row=row, column=col).value
        if val:
            k00_headers.append({'row': row, 'col': col, 'value': str(val)})
result['k00'] = k00_headers

# K.01 Agree SL to GL
ws2 = wb["K.01 Agree SL to GL"]
k01_headers = []
for row in range(1, 15):
    for col in range(1, 20):
        val = ws2.cell(row=row, column=col).value
        if val:
            k01_headers.append({'row': row, 'col': col, 'value': str(val)})
result['k01'] = k01_headers

with open('inspect_headers_result.json', 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

wb.close()
print("Done")
