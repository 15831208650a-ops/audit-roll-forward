#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查C货币资金输出文件
"""

import openpyxl

output_path = r"F:\AI专用\AI工作区\底稿roll forward\output\C SWP 货币资金 20261231 东风鸿远工程咨询有限公司.xlsx"

# data_only=False 查看公式/值
wb = openpyxl.load_workbook(output_path, data_only=False)
ws = wb["C.00 Lead"]

print("=== 主表期初数 ===")
print(f"R38C10 (期末审定数): {ws.cell(row=38, column=10).value}")
print(f"R38C11 (期初审定数): {ws.cell(row=38, column=11).value}")

print("\n=== For Disclosure区域 (R60-R78) ===")
for row in range(60, 79):
    row_data = []
    for col in range(1, 7):
        val = ws.cell(row=row, column=col).value
        if val is not None:
            row_data.append(f"C{col}={val}")
    if row_data:
        print(f"R{row}: {', '.join(row_data)}")

wb.close()
