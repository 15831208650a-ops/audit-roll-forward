import openpyxl

# 验证输出文件
output_path = "output/K1 SWP 固定资产 20261231 六六六有限公司.xlsx"
print(f"=== 验证输出文件: {output_path} ===\n")

wb = openpyxl.load_workbook(output_path, data_only=False)

# 1. 验证 Lead Sheet
ws = wb["K.00 Lead Sheet"]
print("=== K.00 Lead Sheet 验证 ===")
print(f"  B2 (客户名称): {ws.cell(row=2, column=2).value}")
print(f"  C2 (公司名称): {ws.cell(row=2, column=3).value}")
print(f"  B3 (期末): {ws.cell(row=3, column=2).value}")
print(f"  C3 (日期): {ws.cell(row=3, column=3).value}")
print(f"  C4 (分析日期): {ws.cell(row=4, column=3).value}")
print(f"  C5 (TE): {ws.cell(row=5, column=3).value}")
print(f"  C6 (SAD): {ws.cell(row=6, column=3).value}")

# 验证期初数（J列）
print("\n=== 期初数验证（J列，第48行表头后）===")
for row in range(49, 60):
    val = ws.cell(row=row, column=10).value  # J列
    if val is not None:
        print(f"  J{row}: {val}")

# 2. 验证 K.01
ws_k01 = wb["K.01 Agree SL to GL"]
print("\n=== K.01 Agree SL to GL 验证 ===")
print(f"  总行数: {ws_k01.max_row}, 总列数: {ws_k01.max_column}")

# 显示第12行（年初余额）的数据
print(f"\n  第12行数据（年初余额）：")
for col in range(3, 20):
    val = ws_k01.cell(row=12, column=col).value
    if val is not None:
        print(f"    列{col}: {val}")

wb.close()
