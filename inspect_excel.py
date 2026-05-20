import openpyxl
import json

def safe_str(val):
    if val is None:
        return None
    return str(val)

# 检查标准底稿的关键cell
wb = openpyxl.load_workbook('K1 SWP 固定资产 202YMMDD XYZ公司.xlsx', data_only=False)

with open('inspect_result.json', 'w', encoding='utf-8') as f:
    result = {}
    result['template_sheets'] = wb.sheetnames

    # K.00 Lead Sheet
    ws = wb['K.00 Lead Sheet']
    template_k00 = []
    for row in range(1, 10):
        for col in range(1, 10):
            cell = ws.cell(row=row, column=col)
            if cell.value:
                template_k00.append({'row': row, 'col': col, 'value': safe_str(cell.value)})
    result['template_k00'] = template_k00

    # SkywindSettingSheet
    ws2 = wb['SkywindSettingSheet']
    template_sky = []
    for row in range(1, 20):
        for col in range(1, 10):
            cell = ws2.cell(row=row, column=col)
            if cell.value:
                template_sky.append({'row': row, 'col': col, 'value': safe_str(cell.value)})
    result['template_sky'] = template_sky

    # 上年底稿
    wb2 = openpyxl.load_workbook('输入：K1 固定资产 20251231六六六有限公司.xlsx', data_only=False)
    result['prior_sheets'] = wb2.sheetnames

    ws3 = wb2['K.00 Lead Sheet']
    prior_k00 = []
    for row in range(1, 10):
        for col in range(1, 10):
            cell = ws3.cell(row=row, column=col)
            if cell.value:
                prior_k00.append({'row': row, 'col': col, 'value': safe_str(cell.value)})
    result['prior_k00'] = prior_k00

    # 参考资料
    wb3 = openpyxl.load_workbook('输入；参考资料.xlsx', data_only=True)
    result['ref_sheets'] = wb3.sheetnames

    ws4 = wb3['PMTE']
    pmte = []
    for row in range(1, 14):
        for col in range(1, 8):
            cell = ws4.cell(row=row, column=col)
            if cell.value:
                pmte.append({'row': row, 'col': col, 'value': safe_str(cell.value)})
    result['pmte'] = pmte

    json.dump(result, f, ensure_ascii=False, indent=2)
