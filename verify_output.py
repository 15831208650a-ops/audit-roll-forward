import openpyxl

# 验证输出文件
output_path = "output/K1 SWP 固定资产 202YMMDD 六六六有限公司.xlsx"
print(f"=== 验证输出文件: {output_path} ===\n")

wb = openpyxl.load_workbook(output_path, data_only=False)
print(f"Sheets: {wb.sheetnames}\n")

# 1. 验证 Lead Sheet
ws = wb["K.00 Lead Sheet"]
print("=== K.00 Lead Sheet 验证 ===")
print(f"  B2 (客户名称): {ws.cell(row=2, column=2).value}")
print(f"  C2 (公司名称): {ws.cell(row=2, column=3).value}")
print(f"  B3 (期末): {ws.cell(row=3, column=2).value}")
print(f"  C3 (日期): {ws.cell(row=3, column=3).value}")
print(f"  C4 (分析日期): {ws.cell(row=4, column=3).value}")
print(f"  C5 (TE): {ws.cell(row=5, column=3).value}")
print(f"  C6 (SAD): {ws.cell(row=6, column=3).value}")

# 2. 验证期初数（J列）
print("\n=== 期初数验证（J列）===")
for row in range(1, 20):
    cell_value = ws.cell(row=row, column=10).value  # J列
    if cell_value is not None:
        print(f"  J{row}: {cell_value}")

# 3. 验证 K.01
ws_k01 = wb["K.01 Agree SL to GL"]
print("\n=== K.01 Agree SL to GL 验证 ===")
print(f"  总行数: {ws_k01.max_row}, 总列数: {ws_k01.max_column}")

# 查找表头行
header_row = None
for row in range(1, min(30, ws_k01.max_row + 1)):
    for col in range(1, min(ws_k01.max_column + 1, 30)):
        cell_value = ws_k01.cell(row=row, column=col).value
        if cell_value and "期末审定" in str(cell_value):
            header_row = row
            break
    if header_row:
        break

if header_row:
    print(f"\n  表头行: {header_row}")
    print(f"  表头内容:")
    for col in range(1, min(20, ws_k01.max_column + 1)):
        val = ws_k01.cell(row=header_row, column=col).value
        if val:
            print(f"    列{col}: {val}")

    # 显示数据行示例
    print(f"\n  数据行示例 (第{header_row+1}行):")
    for col in range(1, min(20, ws_k01.max_column + 1)):
        val = ws_k01.cell(row=header_row + 1, column=col).value
        if val:
            print(f"    列{col}: {val}")

wb.close()
