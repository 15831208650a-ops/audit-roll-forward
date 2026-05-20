import openpyxl

pmte_path = "输入；参考资料.xlsx"
wb = openpyxl.load_workbook(pmte_path, data_only=True)
ws = wb["CRA"]

print("=== 搜索固定资产相关CRA信息 ===")
for row in range(1, ws.max_row + 1):
    for col in range(1, 12):
        val = ws.cell(row=row, column=col).value
        if val and "固定" in str(val):
            print(f"  ({row},{col}): {val}")

wb.close()
