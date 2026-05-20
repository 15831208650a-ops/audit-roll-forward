#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
详细检查U_exp财务费用子表的结构
"""

import openpyxl

def inspect_sheet(filepath, sheet_name):
    """检查指定Sheet的详细结构"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[sheet_name]

    print(f"\nSheet: {sheet_name}")
    print(f"尺寸: {ws.max_row}x{ws.max_column}")

    # 查找包含"上年"、"期初"、"审定"等字样的单元格
    print("\n查找期初/上年/审定相关单元格:")
    found = False
    for row in range(1, min(ws.max_row + 1, 50)):
        for col in range(1, min(ws.max_column + 1, 20)):
            val = ws.cell(row=row, column=col).value
            if val and any(x in str(val) for x in ["上年", "期初", "审定", "PY", "期末"]):
                print(f"  R{row}C{col}: {str(val)[:80]}")
                found = True

    if not found:
        print("  未找到相关单元格")

    # 读取Row 25-35的内容（预期包含期初数的区域）
    print("\nRow 25-35 内容预览:")
    for row in range(25, min(36, ws.max_row + 1)):
        row_data = []
        for col in range(1, min(20, ws.max_column + 1)):
            val = ws.cell(row=row, column=col).value
            if val and str(val).strip():
                row_data.append(f"C{col}={str(val)[:40]}")
        if row_data:
            print(f"  R{row}: {', '.join(row_data)}")

    wb.close()


if __name__ == "__main__":
    filepath = r"F:\AI专用\AI工作区\底稿roll forward\templates\U_exp SWP other 202YMMDD XYZ公司.xlsx"
    inspect_sheet(filepath, "Uexp财务费用BKD")
