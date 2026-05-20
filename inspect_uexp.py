#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
详细检查U_exp（财务费用）子表结构
"""

import os
import openpyxl

def inspect_uexp(filepath):
    """检查U_exp模板的详细结构"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    result = {"filename": os.path.basename(filepath)}

    # 列出所有sheet
    print(f"文件: {result['filename']}")
    print(f"Sheet列表: {wb.sheetnames}\n")

    # 检查财务费用相关的sheet
    for sheet_name in wb.sheetnames:
        if "财务" in sheet_name or "BKD" in sheet_name or "Lead" in sheet_name:
            ws = wb[sheet_name]
            print(f"\n{'='*60}")
            print(f"[{sheet_name}] 尺寸: {ws.max_row}x{ws.max_column}")

            # 读取前几行内容
            for row in range(1, min(15, ws.max_row + 1)):
                row_data = []
                for col in range(1, min(12, ws.max_column + 1)):
                    val = ws.cell(row=row, column=col).value
                    if val and str(val).strip():
                        row_data.append(f"C{col}={str(val)[:50]}")
                if row_data:
                    print(f"  R{row}: {', '.join(row_data)}")

            # 查找是否有"期初"、"上年"等字样
            print(f"\n  查找期初/上年数据位置:")
            found = False
            for row in range(1, min(80, ws.max_row + 1)):
                for col in range(1, min(20, ws.max_column + 1)):
                    val = ws.cell(row=row, column=col).value
                    if val and any(x in str(val) for x in ["期初", "上年", "年初", "PY", "审定"]):
                        print(f"    R{row}C{col}: {str(val)[:60]}")
                        found = True
            if not found:
                print("    未找到")

    wb.close()
    return result


if __name__ == "__main__":
    inspect_uexp(r"F:\AI专用\AI工作区\底稿roll forward\templates\U_exp SWP other 202YMMDD XYZ公司.xlsx")
