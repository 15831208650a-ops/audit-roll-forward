#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查Lead Sheet表头行的详细列结构
"""

import os
import openpyxl

def inspect_header_row(filepath):
    """找到"期末审定数"所在行，并显示该行的所有列内容"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    result = {"filename": os.path.basename(filepath)}

    lead_sheet = None
    for sheet_name in wb.sheetnames:
        if "Lead" in sheet_name:
            lead_sheet = sheet_name
            break

    if not lead_sheet:
        return result

    ws = wb[lead_sheet]
    result["lead_sheet"] = lead_sheet

    # 查找"期末审定数"所在行
    header_row = None
    for row in range(1, min(80, ws.max_row + 1)):
        for col in range(1, min(20, ws.max_column + 1)):
            val = ws.cell(row=row, column=col).value
            if val and "期末审定数" in str(val):
                header_row = row
                break
        if header_row:
            break

    if not header_row:
        result["error"] = "找不到'期末审定数'"
        return result

    result["header_row"] = header_row

    # 读取该行的所有列内容
    row_data = {}
    for col in range(1, min(20, ws.max_column + 1)):
        val = ws.cell(row=header_row, column=col).value
        if val:
            row_data[f"C{col}"] = str(val)[:30]

    result["columns"] = row_data

    # 再读取下一行（数据行）的内容预览
    data_preview = {}
    for col in range(1, min(20, ws.max_column + 1)):
        val = ws.cell(row=header_row + 1, column=col).value
        if val:
            data_preview[f"C{col}"] = str(val)[:30]
    result["data_preview"] = data_preview

    wb.close()
    return result


def main():
    base_dir = r"F:\AI专用\AI工作区\底稿roll forward"
    files = [f for f in os.listdir(base_dir) if f.endswith('.xlsx') and not f.startswith('~') and not f.startswith('输入')]

    for f in sorted(files):
        filepath = os.path.join(base_dir, f)
        try:
            result = inspect_header_row(filepath)
            print(f"\n{'='*80}")
            print(f"文件: {result['filename']}")
            if "error" in result:
                print(f"  错误: {result['error']}")
                continue

            print(f"  Lead Sheet: {result['lead_sheet']}")
            print(f"  表头行: 第{result['header_row']}行")
            print(f"  表头列内容:")
            for col, text in result['columns'].items():
                print(f"    {col}: {text}")

            print(f"\n  数据行预览(第{result['header_row'] + 1}行):")
            for col, text in list(result['data_preview'].items())[:5]:
                print(f"    {col}: {text}")

        except Exception as e:
            print(f"\n文件 {f} 检查失败: {e}")


if __name__ == "__main__":
    main()
