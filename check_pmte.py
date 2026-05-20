#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查PMTE信息表结构
"""

import openpyxl

def check_pmte(pmte_path):
    wb = openpyxl.load_workbook(pmte_path, data_only=True)

    print(f"Sheets: {wb.sheetnames}")

    if "PMTE" in wb.sheetnames:
        ws = wb["PMTE"]
        print(f"\nPMTE Sheet 前5行数据:")
        for row in range(1, min(6, ws.max_row + 1)):
            row_data = []
            for col in range(1, min(8, ws.max_column + 1)):
                val = ws.cell(row=row, column=col).value
                if val is not None:
                    row_data.append(f"C{col}={val}")
            if row_data:
                print(f"  R{row}: {', '.join(row_data)}")

    wb.close()

check_pmte(r"F:\AI专用\AI工作区\底稿roll forward\pmte_test.xlsx")
