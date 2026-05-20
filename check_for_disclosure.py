#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查For Disclosure表结构
"""

import openpyxl

def check_for_disclosure(prior_path):
    wb = openpyxl.load_workbook(prior_path, data_only=True)
    ws = wb["C.00 Lead"]

    print(f"检查文件: {prior_path}\n")

    # 查找For Disclosure区域
    print("=== For Disclosure区域 ===")
    for row in range(1, ws.max_row + 1):
        val_c2 = ws.cell(row=row, column=2).value
        if val_c2 and "For Disclosure" in str(val_c2):
            print(f"  找到 'For Disclosure' 在第{row}行")
            # 显示前后几行
            for r in range(row - 2, min(row + 10, ws.max_row + 1)):
                row_data = []
                for c in range(1, 8):
                    v = ws.cell(row=r, column=c).value
                    if v is not None:
                        row_data.append(f"C{c}={v}")
                if row_data:
                    print(f"    R{r}: {', '.join(row_data)}")
            break

    wb.close()

if __name__ == "__main__":
    check_for_disclosure(r"F:\AI专用\AI工作区\底稿roll forward\上年底稿文件夹\C SWP 货币资金 20251231 东风鸿远工程咨询有限公司.xlsx")
