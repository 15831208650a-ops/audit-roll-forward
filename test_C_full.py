#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
完整测试C货币资金的Roll Forward
"""

import os
import sys
sys.path.insert(0, r"F:\AI专用\AI工作区\底稿roll forward")

from roll_forward_core import process_single_subject

# 配置
template_path = r"F:\AI专用\AI工作区\底稿roll forward\templates\C SWP 货币资金 202YMMDD XYZ公司.xlsx"
prior_path = r"F:\AI专用\AI工作区\底稿roll forward\上年底稿文件夹\C SWP 货币资金 20251231 东风鸿远工程咨询有限公司.xlsx"
pmte_path = r"F:\AI专用\AI工作区\底稿roll forward\pmte_test.xlsx"
output_dir = r"F:\AI专用\AI工作区\底稿roll forward\output"

subject_config = {
    "code": "C",
    "name": "货币资金",
    "template_file": "C SWP 货币资金 202YMMDD XYZ公司.xlsx",
    "prior_file_pattern": "C*货币资金*{prior_year}*.xlsx",
    "lead_sheet": {
        "sheet_name": "C.00 Lead",
        "header_search_text": "期末审定数",
        "closing_col": 10,
        "opening_col": 11
    },
    "k01": {
        "has_k01": False
    }
}

# 执行处理
success, message, output_path, warnings_list = process_single_subject(
    "C", template_path, prior_path, pmte_path,
    "东风鸿远工程咨询有限公司", "2026-12-31", output_dir, subject_config
)

print(f"处理结果: {success}")
print(f"消息: {message}")
print(f"输出路径: {output_path}")
print(f"警告: {warnings_list}")

# 检查结果
if success and output_path:
    import openpyxl
    wb = openpyxl.load_workbook(output_path, data_only=True)
    ws = wb["C.00 Lead"]

    print("\n=== 检查结果 ===")
    # 1. 检查第38行主表期初数
    print(f"主表R38C10 (期末审定数): {ws.cell(row=38, column=10).value}")
    print(f"主表R38C11 (期初审定数): {ws.cell(row=38, column=11).value}")

    # 2. 检查日期
    print(f"\n日期R3C3: {ws.cell(row=3, column=3).value}")

    # 3. 检查PMTE信息
    print(f"\nTER5C3: {ws.cell(row=5, column=3).value}")
    print(f"SADR6C3: {ws.cell(row=6, column=3).value}")

    # 4. 检查For Disclosure表
    print(f"\nFor Disclosure表:")
    for row in range(71, 78):
        val_c2 = ws.cell(row=row, column=2).value
        val_c3 = ws.cell(row=row, column=3).value
        val_c4 = ws.cell(row=row, column=4).value
        if val_c2:
            print(f"  R{row}C2={val_c2}, C3={val_c3}, C4={val_c4}")

    wb.close()
