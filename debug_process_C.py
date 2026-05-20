#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
调试C货币资金的Roll Forward处理
"""

import os
import shutil
import openpyxl

def debug_process_C():
    template_path = r"F:\AI专用\AI工作区\底稿roll forward\templates\C SWP 货币资金 202YMMDD XYZ公司.xlsx"
    prior_path = r"F:\AI专用\AI工作区\底稿roll forward\上年底稿文件夹\C SWP 货币资金 20251231 东风鸿远工程咨询有限公司.xlsx"
    output_path = r"F:\AI专用\AI工作区\底稿roll forward\output_final\debug_C_output.xlsx"

    # 复制模板
    shutil.copy2(template_path, output_path)
    os.chmod(output_path, 0o666)

    # 打开文件
    wb_new = openpyxl.load_workbook(output_path)
    wb_prior_values = openpyxl.load_workbook(prior_path, data_only=True)

    ws_new = wb_new["C.00 Lead"]
    ws_prior = wb_prior_values["C.00 Lead"]

    # 查找表头行
    header_row = None
    for row in range(1, 80):
        for col in range(1, 20):
            val = ws_new.cell(row=row, column=col).value
            if val and "期末审定数" in str(val):
                header_row = row
                print(f"找到'期末审定数'在第{row}行, 第{col}列")
                break
        if header_row:
            break

    closing_col = 10
    opening_col = 11

    print(f"\n表头行: {header_row}")
    print(f"期末审定数列: C{closing_col}, 期初审定数列: C{opening_col}")

    # 检查模板和新文件中C38的值（处理前）
    print(f"\n=== 处理前 ===")
    print(f"新文件C38C10 (期末审定数): {ws_new.cell(row=38, column=10).value}")
    print(f"新文件C38C11 (期初审定数): {ws_new.cell(row=38, column=11).value}")
    print(f"旧文件C38C10 (期末审定数): {ws_prior.cell(row=38, column=10).value}")
    print(f"旧文件C38C11 (期初审定数): {ws_prior.cell(row=38, column=11).value}")

    # 执行复制
    copied = 0
    for row in range(header_row + 1, ws_prior.max_row + 1):
        prior_cell = ws_prior.cell(row=row, column=closing_col)
        new_cell = ws_new.cell(row=row, column=opening_col)

        if prior_cell.value is not None:
            print(f"  复制 R{row}: C{closing_col}={prior_cell.value} -> C{opening_col}")
            new_cell.value = prior_cell.value
            copied += 1

    print(f"\n=== 处理后 ===")
    print(f"新文件C38C10 (期末审定数): {ws_new.cell(row=38, column=10).value}")
    print(f"新文件C38C11 (期初审定数): {ws_new.cell(row=38, column=11).value}")

    # 保存
    wb_new.save(output_path)
    print(f"\n已保存到: {output_path}")

    wb_new.close()
    wb_prior_values.close()

if __name__ == "__main__":
    debug_process_C()
