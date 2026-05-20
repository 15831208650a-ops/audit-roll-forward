#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
调试For Disclosure处理2
"""

import openpyxl

# 模拟process_lead_sheet中的For Disclosure处理逻辑
prior_path = r"F:\AI专用\AI工作区\底稿roll forward\上年底稿文件夹\C SWP 货币资金 20251231 东风鸿远工程咨询有限公司.xlsx"
new_path = r"F:\AI专用\AI工作区\底稿roll forward\templates\C SWP 货币资金 202YMMDD XYZ公司.xlsx"

# 旧文件 - data_only=True
wb_prior = openpyxl.load_workbook(prior_path, data_only=True)
ws_prior = wb_prior["C.00 Lead"]

# 新文件 - data_only=False
wb_new = openpyxl.load_workbook(new_path, data_only=False)
ws_new = wb_new["C.00 Lead"]

# 1. 在旧文件中找到Disclosure
print("=== 旧文件Disclosure ===")
disclosure_start_row = None
for row in range(1, ws_prior.max_row + 1):
    for col in range(1, 5):
        val = ws_prior.cell(row=row, column=col).value
        if val and "Disclosure" in str(val):
            disclosure_start_row = row
            print(f"找到Disclosure在R{row}C{col}: {val}")
            break
    if disclosure_start_row:
        break

# 2. 在旧文件中找到表头行
fd_prior_header_row = None
fd_prior_closing_col = None
fd_prior_opening_col = None

for row in range(disclosure_start_row, min(disclosure_start_row + 10, ws_prior.max_row + 1)):
    for col in range(1, 20):
        val = ws_prior.cell(row=row, column=col).value
        if val and "期末审定数" in str(val):
            fd_prior_header_row = row
            fd_prior_closing_col = col
            for next_col in range(col + 1, min(col + 5, ws_prior.max_column + 1)):
                next_val = ws_prior.cell(row=row, column=next_col).value
                if next_val and "期初" in str(next_val):
                    fd_prior_opening_col = next_col
                    break
            break
    if fd_prior_header_row:
        break

print(f"旧文件表头行: R{fd_prior_header_row}, 期末列: C{fd_prior_closing_col}, 期初列: C{fd_prior_opening_col}")

# 3. 在新文件中找到For Disclosure表头行
fd_new_header_row = None
fd_new_opening_col = None

for row in range(1, ws_new.max_row + 1):
    val = ws_new.cell(row=row, column=2).value
    if val and "项目名称" in str(val):
        found_disclosure = False
        for check_row in range(max(1, row - 5), row):
            for check_col in range(1, 5):
                check_val = ws_new.cell(row=check_row, column=check_col).value
                if check_val and "Disclosure" in str(check_val):
                    found_disclosure = True
                    break
            if found_disclosure:
                break
        if found_disclosure:
            fd_new_header_row = row
            for col in range(1, 20):
                col_val = ws_new.cell(row=row, column=col).value
                if col_val and "期初" in str(col_val):
                    fd_new_opening_col = col
                    break
            if not fd_new_opening_col:
                for col in range(1, 20):
                    col_val = ws_new.cell(row=row, column=col).value
                    if col_val and "期末审定数" in str(col_val):
                        fd_new_opening_col = col + 1
                        break
            if not fd_new_opening_col:
                fd_new_opening_col = 4
            break

print(f"新文件表头行: R{fd_new_header_row}, 期初列: C{fd_new_opening_col}")

# 4. 构建旧文件的项目名称字典
prior_items = {}
for row in range(fd_prior_header_row + 1, ws_prior.max_row + 1):
    item_name = ws_prior.cell(row=row, column=2).value
    if item_name and str(item_name).strip():
        val = ws_prior.cell(row=row, column=fd_prior_closing_col).value
        if val is not None:
            prior_items[str(item_name).strip()] = val
            print(f"  旧文件项目: '{item_name}' = {val}")

# 5. 在新文件中查找匹配
print(f"\n=== 新文件For Disclosure数据行 ===")
for row in range(fd_new_header_row + 1, ws_new.max_row + 1):
    item_name = ws_new.cell(row=row, column=2).value
    if item_name:
        print(f"  R{row}: '{item_name}'")
        if str(item_name).strip() in prior_items:
            print(f"    -> 匹配成功! 值={prior_items[str(item_name).strip()]}")
        else:
            print(f"    -> 未匹配")

wb_prior.close()
wb_new.close()
