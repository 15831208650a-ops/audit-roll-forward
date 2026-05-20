#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
详细检查各科目模板的结构差异
"""

import os
import json
import openpyxl

def inspect_file(filepath):
    """详细检查单个文件"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    result = {"filename": os.path.basename(filepath), "sheets": {}}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_info = {
            "max_row": ws.max_row,
            "max_col": ws.max_column
        }

        # 查找Lead Sheet中的关键行
        if "Lead" in sheet_name:
            for row in range(1, min(80, ws.max_row + 1)):
                for col in range(1, 15):
                    val = ws.cell(row=row, column=col).value
                    if val and "期末审定数" in str(val):
                        sheet_info["期末审定数"] = {"row": row, "col": col}
                        break
                if "期末审定数" in sheet_info:
                    break

            # 读取Lead Sheet前几行内容（公司名、日期等位置）
            header_info = {}
            for row in range(1, 10):
                row_data = []
                for col in range(1, 6):
                    val = ws.cell(row=row, column=col).value
                    if val:
                        row_data.append(f"C{col}={str(val)[:30]}")
                if row_data:
                    header_info[f"R{row}"] = row_data
            sheet_info["header_preview"] = header_info

        # 查找K.01或类似Sheet的结构
        if any(x in sheet_name for x in ["Agree", "BKD", "01"]):
            # 查找"年初余额"
            for row in range(1, min(80, ws.max_row + 1)):
                for col in range(1, 5):
                    val = ws.cell(row=row, column=col).value
                    if val and "年初余额" in str(val):
                        sheet_info["年初余额"] = {"row": row, "col": col}
                        break
                if "年初余额" in sheet_info:
                    break

            # 读取Row 8-15的关键列头信息
            header_data = {}
            for check_row in range(8, 16):
                row_data = []
                for col in range(1, min(20, ws.max_column + 1)):
                    val = ws.cell(row=check_row, column=col).value
                    if val and str(val).strip():
                        row_data.append(f"C{col}={str(val)[:20]}")
                if row_data:
                    header_data[f"R{check_row}"] = row_data
            sheet_info["row_headers"] = header_data

        result["sheets"][sheet_name] = sheet_info

    wb.close()
    return result


def main():
    base_dir = r"F:\AI专用\AI工作区\底稿roll forward"
    files = [f for f in os.listdir(base_dir) if f.endswith('.xlsx') and not f.startswith('~') and not f.startswith('输入')]

    all_results = []
    for f in sorted(files):
        try:
            result = inspect_file(os.path.join(base_dir, f))
            all_results.append(result)
        except Exception as e:
            all_results.append({"filename": f, "error": str(e)})

    # 写入JSON文件
    output_path = os.path.join(base_dir, "inspect_result.json")
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(all_results, f, ensure_ascii=False, indent=2)

    print(f"结果已保存到: {output_path}")

    # 打印汇总
    for r in all_results:
        if 'error' in r:
            print(f"\n{'='*60}")
            print(f"文件: {r['filename']} - 错误: {r['error']}")
            continue

        print(f"\n{'='*60}")
        print(f"文件: {r['filename']}")
        for sheet, info in r['sheets'].items():
            print(f"  [{sheet}] {info['max_row']}x{info['max_col']}")
            if "期末审定数" in info:
                p = info["期末审定数"]
                print(f"    '期末审定数' 行{p['row']} 列{p['col']}")
            if "年初余额" in info:
                p = info["年初余额"]
                print(f"    '年初余额' 行{p['row']} 列{p['col']}")
            if "row_headers" in info:
                print(f"    行头预览: {str(info['row_headers'])[:150]}")


if __name__ == "__main__":
    main()
