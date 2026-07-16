import json
import math
import os
import posixpath
import sys
import zipfile
from copy import copy
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook

from cra_support import parse_cra_paste_text
from roll_forward_core import (
    RollForwardWarnings,
    SubjectConfig,
    find_c_bkd_structure,
    find_expense_bkd_header_row,
    find_expense_bkd_total_row,
    find_header_col,
    find_header_col_near,
    find_header_row,
    find_l2_business_end_col,
    find_label_cell,
    find_prior_file,
    find_q1_bkd_header_row,
    normalize_text,
    process_c_bkd_basic_info,
    process_c_cutoff_wording,
    process_j1_cip_long_aging,
    process_lead_sheet,
)


ROOT = Path(__file__).resolve().parent.parent
SOURCE = Path(__file__).resolve().parent
TEMPLATES = ROOT / "templates"
PRIOR_DIR = ROOT / "常规样本" / "吉安"
OUTPUT_DIR = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "tmp_live_integrated_corefix5_final"

FAILURES = []
PASSES = []


def check(condition, message):
    if condition:
        PASSES.append(message)
    else:
        FAILURES.append(message)


def workbook_path(prefix):
    matches = sorted(OUTPUT_DIR.glob(f"{prefix}*.xlsx"))
    if not matches:
        raise FileNotFoundError(f"Missing output workbook: {prefix}")
    return matches[-1]


def sheet_name(wb, prefix):
    return next(name for name in wb.sheetnames if name.startswith(prefix))


def is_roll_highlight(cell):
    return (
        cell.fill.fill_type == "solid"
        and cell.fill.fgColor.type == "rgb"
        and str(cell.fill.fgColor.rgb).upper().endswith("FFFF99")
    )


def has_overlapping_merges(ws):
    ranges = list(ws.merged_cells.ranges)
    for index, first in enumerate(ranges):
        for second in ranges[index + 1 :]:
            separated = (
                first.max_row < second.min_row
                or second.max_row < first.min_row
                or first.max_col < second.min_col
                or second.max_col < first.min_col
            )
            if not separated:
                return True
    return False


def relationship_source(rels_name):
    if rels_name == "_rels/.rels":
        return ""
    return posixpath.join(
        posixpath.dirname(posixpath.dirname(rels_name)),
        posixpath.basename(rels_name)[:-5],
    )


def check_xlsx_package(path):
    errors = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if archive.testzip() is not None:
            errors.append("CRC failure")
        for name in names:
            if not name.endswith((".xml", ".rels")):
                continue
            try:
                root = ElementTree.fromstring(archive.read(name))
            except Exception as exc:
                errors.append(f"Invalid XML {name}: {exc}")
                continue
            if not name.endswith(".rels"):
                continue
            source = relationship_source(name)
            base = posixpath.dirname(source)
            for rel in root:
                if rel.attrib.get("TargetMode") == "External":
                    continue
                target = rel.attrib.get("Target", "")
                resolved = (
                    target.lstrip("/")
                    if target.startswith("/")
                    else posixpath.normpath(posixpath.join(base, target))
                )
                if resolved not in names:
                    errors.append(f"Missing relationship target {name} -> {resolved}")
    wb = load_workbook(path, data_only=False)
    for ws in wb.worksheets:
        if has_overlapping_merges(ws):
            errors.append(f"Overlapping merges: {ws.title}")
        for validation in ws.data_validations.dataValidation:
            try:
                str(validation.sqref)
            except Exception as exc:
                errors.append(f"Invalid validation range {ws.title}: {exc}")
    wb.close()
    check(not errors, f"{path.name}: ZIP/XML/relationships/merge ranges clean")
    if errors:
        FAILURES.extend(f"{path.name}: {error}" for error in errors)


def numeric_equal(left, right):
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), rel_tol=1e-9, abs_tol=0.01)
    return left == right


def expense_records(ws, value_kind):
    header = find_expense_bkd_header_row(ws)
    total = find_expense_bkd_total_row(ws, header)
    if value_kind == "current":
        value_col = find_header_col(ws, header, ["本期账面审定数", "本期审定数", "本期数"])
    else:
        value_col = find_header_col(ws, header, ["上期末审定数", "上期审定数", "上年数", "PY"])
    if not value_col:
        keywords = ["本期账面审定数", "本期审定数"] if value_kind == "current" else ["上期末审定数", "上期审定数", "PY"]
        value_col = find_header_col_near(ws, header, keywords, row_offsets=(-1, 0))
    descriptor_cols = []
    for keywords in (["账套名称/账套编码", "账套名称"], ["科目编码"], ["科目名称"]):
        col = find_header_col(ws, header, keywords)
        if col:
            descriptor_cols.append(col)
    records = []
    for row in range(header + 1, total):
        descriptor = tuple(normalize_text(ws.cell(row, col).value) for col in descriptor_cols)
        if any(descriptor):
            records.append((descriptor, ws.cell(row, value_col).value))
    return records


def check_expense_roll(prior_path, output_path, sheet_prefix):
    prior_wb = load_workbook(prior_path, data_only=True)
    output_wb = load_workbook(output_path, data_only=False)
    prior_records = expense_records(prior_wb[sheet_name(prior_wb, sheet_prefix)], "current")
    output_records = expense_records(output_wb[sheet_name(output_wb, sheet_prefix)], "py")
    prior_wb.close()
    output_wb.close()
    same = len(prior_records) == len(output_records) and all(
        p_desc == o_desc and numeric_equal(p_value, o_value)
        for (p_desc, p_value), (o_desc, o_value) in zip(prior_records, output_records)
    )
    check(same, f"{sheet_prefix}: prior current values rolled to current PY without row loss")


def check_q1(path):
    wb = load_workbook(path, data_only=False)
    q101 = wb[sheet_name(wb, "Q1.01")]
    header = find_q1_bkd_header_row(q101)
    total = next(
        row
        for row in range(header + 1, q101.max_row + 1)
        if any(normalize_text(q101.cell(row, col).value) == "合计" for col in range(1, 8))
    )
    right_side_ok = all(
        q101.cell(row, 15).value in (None, "")
        and q101.cell(row, 17).value in (None, "")
        and q101.cell(row, 18).value in (None, "")
        and isinstance(q101.cell(row, 16).value, str)
        and q101.cell(row, 16).value.startswith("=")
        for row in range(header + 1, total)
    )
    check(right_side_ok, "Q1.01: current/non-current/covenant/Notes1 prior values are not rolled")
    for prefix in ("Q1.02b", "Q1.02c"):
        ws = wb[sheet_name(wb, prefix)]
        count = sum(is_roll_highlight(cell) for row in ws.iter_rows() for cell in row)
        check(count == 0, f"{prefix}: no roll-forward yellow added")
    q105 = wb[sheet_name(wb, "Q1.05")]
    fixed_labels = ("B37", "B38", "B39", "B41", "B43", "B48", "B51", "B55", "B59", "B60", "B61", "B63", "B65", "B70", "B73", "B77")
    check(all(not is_roll_highlight(q105[cell]) for cell in fixed_labels), "Q1.05: template labels are not marked yellow")
    check(not has_overlapping_merges(q105), "Q1.05: no overlapping merged cells")
    template_wb = load_workbook(TEMPLATES / "Q1 SWP 银行借款 202YMMDD XYZ公司.xlsx", data_only=False)
    template_q105 = template_wb[sheet_name(template_wb, "Q1.05")]
    check(
        all(q105.column_dimensions[col].width == template_q105.column_dimensions[col].width for col in "BCDEFGHI"),
        "Q1.05: current-template column widths retained",
    )
    template_wb.close()
    wb.close()

    drawing_rows = []
    media = []
    ns = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}
    with zipfile.ZipFile(path) as archive:
        media = [name for name in archive.namelist() if name.startswith("xl/media/")]
        for name in archive.namelist():
            if not name.startswith("xl/drawings/drawing") or not name.endswith(".xml"):
                continue
            root = ElementTree.fromstring(archive.read(name))
            for anchor in list(root):
                marker = anchor.find("xdr:from", ns)
                if marker is not None and anchor.find("xdr:pic", ns) is not None:
                    drawing_rows.append(int(marker.find("xdr:row", ns).text))
    check(len(media) == 7, "Q1: exactly two Note images and five Q1.05 evidence images copied")
    check(sorted(drawing_rows) == [34, 34, 42, 42, 42, 64, 64], "Q1: copied images follow rebuilt Note blocks")


def first_body_value(ws, anchor_row, max_rows=8, col_start=1, col_end=30):
    for row in range(anchor_row + 1, min(ws.max_row, anchor_row + max_rows) + 1):
        for col in range(col_start, min(ws.max_column, col_end) + 1):
            value = ws.cell(row, col).value
            if value not in (None, "") and not (isinstance(value, str) and value.startswith("=")):
                return row, col, value
    return None, None, None


def nearest_adjacent_value(ws, row, label_col):
    for distance in range(1, min(ws.max_column, 12) + 1):
        for col in (label_col - distance, label_col + distance):
            if col < 1 or col > ws.max_column:
                continue
            value = ws.cell(row, col).value
            if value not in (None, "") and not (isinstance(value, str) and value.startswith("=")):
                return row, col, value
    return None, None, None


def c_bkd_records(ws):
    header, columns, marker = find_c_bkd_structure(ws)
    records = []
    for row in range(header + 1, marker):
        record = tuple(ws.cell(row, columns[field]).value for field in columns)
        if any(value not in (None, "") for value in record):
            records.append((row, record))
    return columns, records


def check_c_roll(path):
    prior_path = PRIOR_DIR / "C SWP 货币资金 20251231 吉安三强.xlsx"
    prior = load_workbook(prior_path, data_only=False)
    prior_values = load_workbook(prior_path, data_only=True)
    output = load_workbook(path, data_only=False)
    template = load_workbook(TEMPLATES / "C SWP 货币资金 202YMMDD XYZ公司.xlsx", data_only=False)

    source_cols, source_records = c_bkd_records(prior_values["C.00 BKD"])
    target_cols, target_records = c_bkd_records(output["C.00 BKD"])
    check(
        [record for _, record in source_records] == [record for _, record in target_records],
        "C.00 BKD: all six account identity fields roll for every real account row",
    )
    check(len(target_records) == 14, "C.00 BKD: template expands to all 14 Ji'an account rows")
    check(
        all(
            not is_roll_highlight(output["C.00 BKD"].cell(row, target_cols[field]))
            for row, _ in target_records
            for field in target_cols
        ),
        "C.00 BKD: always-rolled identity fields are not wording-highlighted",
    )

    prior_cutoff = prior["C.03 Cutoff"]
    output_cutoff = output["C.03 Cutoff"]
    source_period_label = find_label_cell(prior_cutoff, ("使用的截止期间",))
    target_period_label = find_label_cell(output_cutoff, ("使用的截止期间",))
    source_period = nearest_adjacent_value(prior_cutoff, *source_period_label)
    target_period = nearest_adjacent_value(output_cutoff, *target_period_label)
    check(source_period[2] == target_period[2], "C.03 Cutoff: labeled cutoff period rolls dynamically")
    check(is_roll_highlight(output_cutoff.cell(target_period[0], target_period[1])), "C.03 Cutoff: cutoff period is highlighted")

    source_reason_label = find_label_cell(prior_cutoff, ("所使用截止期间的理由",))
    target_reason_label = find_label_cell(output_cutoff, ("所使用截止期间的理由",))
    source_reason = first_body_value(prior_cutoff, source_reason_label[0])
    target_reason = first_body_value(output_cutoff, target_reason_label[0])
    check(source_reason[2] == target_reason[2], "C.03 Cutoff: labeled rationale rolls dynamically")
    check(is_roll_highlight(output_cutoff.cell(target_reason[0], target_reason[1])), "C.03 Cutoff: rationale is highlighted")

    fixture = template
    process_c_bkd_basic_info(prior_values["C.00 BKD"], fixture["C.00 BKD"])
    fixture_period_label = find_label_cell(fixture["C.03 Cutoff"], ("使用的截止期间",))
    fixture_period = nearest_adjacent_value(fixture["C.03 Cutoff"], *fixture_period_label)
    check(fixture_period[2] is None, "C fixture: BKD identity roll is independent of wording")
    process_c_cutoff_wording(prior["C.03 Cutoff"], prior_values["C.03 Cutoff"], fixture["C.03 Cutoff"])
    fixture_period = nearest_adjacent_value(fixture["C.03 Cutoff"], *fixture_period_label)
    check(fixture_period[2] == source_period[2], "C fixture: Cutoff wording rolls only when its wording step runs")

    prior.close()
    prior_values.close()
    output.close()
    template.close()


def check_j1_wording(path):
    prior_path = PRIOR_DIR / "J1 SWP 在建工程 20251231吉安.xlsx"
    prior = load_workbook(prior_path, data_only=False)
    prior_values = load_workbook(prior_path, data_only=True)
    output = load_workbook(path, data_only=False)
    template = load_workbook(TEMPLATES / "J1 SWP 在建工程 202YMMDD XYZ公司.xlsx", data_only=False)
    prior_ws = prior["J.03 CIP长期挂账"]
    prior_values_ws = prior_values["J.03 CIP长期挂账"]
    output_ws = output["J.03 CIP长期挂账"]
    template_ws = template["J.03 CIP长期挂账"]

    source_reason_anchor = find_label_cell(prior_ws, ("选择待测试项目的理由",))[0]
    target_reason_anchor = find_label_cell(output_ws, ("选择待测试项目的理由",))[0]
    source_reason = first_body_value(prior_ws, source_reason_anchor)
    target_reason = first_body_value(output_ws, target_reason_anchor)
    check(source_reason[2] == target_reason[2], "J.03: selection rationale rolls by its prompt")
    check(is_roll_highlight(output_ws.cell(target_reason[0], target_reason[1])), "J.03: selection rationale is highlighted")

    note_anchor = find_label_cell(prior_ws, ("标记图例",))[0]
    fixture_text = "Fixture J.03 Note"
    prior_ws.cell(note_anchor + 1, 2).value = fixture_text
    prior_values_ws.cell(note_anchor + 1, 2).value = fixture_text
    process_j1_cip_long_aging(prior_ws, prior_values_ws, template_ws)
    target_note_anchor = find_label_cell(template_ws, ("标记图例",))[0]
    target_note = first_body_value(template_ws, target_note_anchor)
    check(target_note[2] == fixture_text, "J.03 fixture: populated Note rolls by prompt")
    check(is_roll_highlight(template_ws.cell(target_note[0], target_note[1])), "J.03 fixture: populated Note is highlighted")

    prior.close()
    prior_values.close()
    output.close()
    template.close()


def l2_note_values(ws):
    values = []
    for row in range(1, ws.max_row + 1):
        if normalize_text(ws.cell(row, 2).value).lower().startswith("notes"):
            values.append(first_body_value(ws, row, max_rows=4, col_start=3, col_end=8))
    return values


def check_l2_full_width_and_wording(path):
    prior = load_workbook(PRIOR_DIR / "L2 SWP 长期待摊费用 20251231吉安.xlsx", data_only=False)
    output = load_workbook(path, data_only=False)
    ws = output["L2.01.1 BKD"]
    header = find_header_row(ws, "项目编码", (1, 80))
    total = next(
        row
        for row in range(header + 1, ws.max_row + 1)
        if any(normalize_text(ws.cell(row, col).value) == "合计" for col in range(1, 10))
    )
    business_end = find_l2_business_end_col(ws, header)
    check(business_end == 39, "L2.01.1: real business boundary reaches AM and excludes trailing plug-in columns")
    check(total - header - 1 == 14, "L2.01.1: all 14 Ji'an detail rows fit before total")

    formula_cols = (20, 26, 27, 28, 30, 31, 32, 34, 36, 37, 38)
    check(
        all(
            isinstance(ws.cell(row, col).value, str) and ws.cell(row, col).value.startswith("=")
            for row in range(header + 1, total)
            for col in formula_cols
        ),
        "L2.01.1: every expanded row keeps formulas through the downstream business sections",
    )
    check(
        all(getattr(ws.cell(row, 35).value, "ref", None) == f"AI{row}" for row in range(header + 1, total)),
        "L2.01.1: every expanded row has a valid row-local AI array formula",
    )
    check(
        all(ws.cell(header + 1, col).style_id == ws.cell(total - 1, col).style_id for col in range(19, business_end + 1)),
        "L2.01.1: downstream styles extend through the last detail row",
    )
    check(
        any(
            any(cell_range.min_col <= 33 <= cell_range.max_col and cell_range.min_row <= header + 1 and cell_range.max_row >= total - 1 for cell_range in validation.sqref.ranges)
            for validation in ws.data_validations.dataValidation
        ),
        "L2.01.1: AG sample-type validation covers every expanded row",
    )
    check(
        ws.cell(total, 19).value == f"=SUM(S{header + 1}:S{total - 1})"
        and ws.cell(total, 30).value == f"=SUM(AD{header + 1}:AD{total - 1})"
        and ws.cell(total, 36).value == f"=SUM(AJ{header + 1}:AJ{total - 1})",
        "L2.01.1: downstream total formulas use the full expanded detail range",
    )
    downstream_formula = next(
        ws.cell(row, 7).value
        for row in range(total + 1, ws.max_row + 1)
        if isinstance(ws.cell(row, 7).value, str) and ws.cell(row, 7).value.startswith("=SUM(")
    )
    check(f"-P{total}" in downstream_formula, "L2.01.1: table 2 reconciliation formula follows the moved table total")

    prior_notes = l2_note_values(prior["L2.01.1 BKD"])
    output_notes = l2_note_values(ws)
    check(
        len(prior_notes) == len(output_notes) == 3
        and [item[2] for item in prior_notes] == [item[2] for item in output_notes],
        "L2.01.1: all three real Notes roll to their corresponding boxes",
    )
    check(
        all(is_roll_highlight(ws.cell(row, col)) for row, col, value in output_notes if value not in (None, "")),
        "L2.01.1: every rolled Note response is highlighted",
    )

    prior_lead = prior["L2.00 Lead"]
    output_lead = output["L2.00 Lead"]
    prior_adjustment = find_label_cell(prior_lead, ("调整汇总表",))[0]
    output_adjustment = find_label_cell(output_lead, ("调整汇总表",))[0]
    prior_value = first_body_value(prior_lead, prior_adjustment)
    output_value = first_body_value(output_lead, output_adjustment)
    check(prior_value[2] == output_value[2], "L2.00 Lead: populated adjustment summary rolls when wording is enabled")
    check(is_roll_highlight(output_lead.cell(output_value[0], output_value[1])), "L2.00 Lead: adjustment content is highlighted")
    check(
        all(not is_roll_highlight(output_lead.cell(output_adjustment, col)) for col in range(1, 12)),
        "L2.00 Lead: fixed adjustment-summary labels are not highlighted",
    )

    prior.close()
    output.close()


def thin_outer_border(ws, cell_range):
    cells = ws[cell_range]
    return (
        all(cell.border.top.style == "thin" for cell in cells[0])
        and all(cell.border.bottom.style == "thin" for cell in cells[-1])
        and all(row[0].border.left.style == "thin" for row in cells)
        and all(row[-1].border.right.style == "thin" for row in cells)
    )


def check_uexp(path):
    wb = load_workbook(path, data_only=False)
    lead = wb["Uexp_Lead"]
    bkd = wb["Uexp财务费用BKD"]
    for ws, ranges in (
        (lead, ("C13:G15", "C35:G37")),
        (bkd, ("C19:G21", "C46:G46", "C47:G49", "C50:G51", "C52:G52")),
    ):
        merged = {str(item) for item in ws.merged_cells.ranges}
        for cell_range in ranges:
            check(cell_range in merged, f"Uexp {ws.title} {cell_range}: merged wording box retained")
            check(thin_outer_border(ws, cell_range), f"Uexp {ws.title} {cell_range}: thin continuous outer border")
            check(ws[cell_range.split(":")[0]].alignment.wrap_text is True, f"Uexp {ws.title} {cell_range}: wrap text enabled")
    wb.close()


def check_l2_expansion_fixture():
    template_path = TEMPLATES / "L2 SWP 长期待摊费用 202YMMDD XYZ公司.xlsx"
    prior_path = PRIOR_DIR / "L2 SWP 长期待摊费用 20251231吉安.xlsx"
    prior_formula = load_workbook(prior_path, data_only=False)
    prior_values = load_workbook(prior_path, data_only=True)
    output_wb = load_workbook(template_path, data_only=False)
    for workbook in (prior_formula, prior_values):
        ws = workbook["L2.00 Lead"]
        header = find_header_row(ws, "期末审定数", (1, 80))
        closing_col = find_header_col(ws, header, ["期末审定数"])
        ws.insert_rows(header + 3, 5)
        for offset in range(5):
            row = header + 3 + offset
            ws.cell(row, 2, f"Fixture company {offset + 3}")
            ws.cell(row, 3, f"18{offset + 3:02d}")
            ws.cell(row, 4, f"Fixture item {offset + 3}")
            ws.cell(row, closing_col, 1000 + offset)
    config = SubjectConfig(SOURCE / "subjects_config.json").get_subject("L2")["lead_sheet"]
    lead_warnings = RollForwardWarnings()
    process_lead_sheet(
        prior_formula["L2.00 Lead"],
        output_wb["L2.00 Lead"],
        prior_values["L2.00 Lead"],
        {"PM": 1, "TE": 1, "SAD": 1},
        "2026-06-30",
        "Fixture company",
        config,
        lead_warnings,
    )
    ws = output_wb["L2.00 Lead"]
    header = find_header_row(ws, "期末审定数", (1, 80))
    total = next(
        row
        for row in range(header + 1, ws.max_row + 1)
        if any(normalize_text(ws.cell(row, col).value) == "合计" for col in range(1, 8))
    )
    detail_start = header + 1
    detail_end = detail_start + 7 - 1
    check(total == detail_end + 1, "L2 fixture: Lead total row moves below seven detail rows")
    check(
        all(ws.cell(row, 10).value == f"=G{row}+I{row}" and ws.cell(row, 13).value == f"=J{row}+L{row}" for row in range(detail_start, detail_end + 1)),
        "L2 fixture: every expanded detail row retains J/M formulas",
    )
    check(ws.cell(total, 7).value == f"=SUM(G{detail_start}:G{detail_end})", "L2 fixture: total formula uses expanded detail range")
    below_text = [normalize_text(ws.cell(row, col).value) for row in range(total + 1, min(total + 12, ws.max_row) + 1) for col in range(1, 14)]
    check("Rx" in below_text and "A3" in below_text and "Diff" in below_text, "L2 fixture: Rx/A3/Diff area moves below total")
    check(
        not any("未找到Level数据" in warning or "未找到RP数据" in warning for warning in lead_warnings),
        "Warnings: optional PMTE Level/RP absence is silent",
    )
    prior_formula.close()
    prior_values.close()
    output_wb.close()


def check_repair_roots(l2_path, vcvd_path):
    with zipfile.ZipFile(vcvd_path) as archive:
        external_link_parts = [name for name in archive.namelist() if "externalLinks/" in name]
    check(not external_link_parts, "VCVD: unstable external-link cache is not retained")

    wb = load_workbook(l2_path, data_only=False)
    ws = wb["L2.01.1 BKD"]
    header = find_header_row(ws, "项目编码", (1, 80))
    total = next(
        row
        for row in range(header + 1, ws.max_row + 1)
        if any(normalize_text(ws.cell(row, col).value) == "合计" for col in range(1, 8))
    )
    business_end = find_l2_business_end_col(ws, header)
    plug_in_formula_rows = [
        row
        for row in range(header + 1, total)
        if any(
            isinstance(ws.cell(row, col).value, str) and ws.cell(row, col).value.startswith("=")
            for col in range(business_end + 1, ws.max_column + 1)
        )
    ]
    check(not plug_in_formula_rows, "L2.01.1: expanded records stop before trailing plug-in columns")
    check(ws.cell(total, 7).value == f"=SUM(G{header + 1}:G{total - 1})", "L2.01.1: total starts at first detail row")
    wb.close()


def check_cra_applicable_variants():
    header = "科目\t认定\tCRA\t比例\t是否适用"
    for raw in ("N", "N/A", "n/a", "no", "否", "不适用"):
        records = parse_cra_paste_text(
            header + f"\n管理费用\t发生\tModerate\t50%\t{raw}",
            selected_subjects=["UexpVCVD"],
        )
        record = records[0]
        check(record["applicable"] is False and record["cra_level"] == "N/A" and record["ratio_text"] == "N/A", f"CRA applicable={raw}: writes N/A")
    for text, label in (
        (header + "\n管理费用\t发生\tModerate\t50%\t", "blank applicable"),
        ("科目\t认定\tCRA\t比例\n管理费用\t发生\tModerate\t50%", "missing applicable column"),
    ):
        record = parse_cra_paste_text(text, selected_subjects=["UexpVCVD"])[0]
        check(record["applicable"] is True and record["cra_level"] == "Moderate", f"CRA {label}: risk level retained")


def main():
    config = SubjectConfig(SOURCE / "subjects_config.json")
    matched_vcvd = find_prior_file(PRIOR_DIR, "UexpVCVD", "2025", config.get_subject("UexpVCVD"))
    matched_uexp = find_prior_file(PRIOR_DIR, "Uexp", "2025", config.get_subject("Uexp"))
    check(matched_vcvd and "VC&VD" in Path(matched_vcvd).name, "VCVD matcher selects the VC&VD prior workbook")
    check(matched_uexp and "财务费用" in Path(matched_uexp).name, "Uexp matcher selects the finance expense prior workbook")

    c = workbook_path("C SWP")
    j1 = workbook_path("J1 SWP")
    q1 = workbook_path("Q1 SWP")
    vcvd = workbook_path("U_exp SWP VC&VD")
    l2 = workbook_path("L2 SWP")
    uexp = workbook_path("U_exp SWP other")
    for path in (c, j1, q1, vcvd, l2, uexp):
        check_xlsx_package(path)

    check_c_roll(c)
    check_j1_wording(j1)
    check_q1(q1)
    check_expense_roll(Path(matched_vcvd), vcvd, "VC.00")
    check_expense_roll(Path(matched_vcvd), vcvd, "VD.00")
    check_expense_roll(Path(matched_uexp), uexp, "Uexp财务费用BKD")
    check_uexp(uexp)
    check_l2_expansion_fixture()
    check_l2_full_width_and_wording(l2)
    check_repair_roots(l2, vcvd)
    check_cra_applicable_variants()

    report = {"passes": len(PASSES), "failures": FAILURES, "checked_output": str(OUTPUT_DIR)}
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if FAILURES:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
