#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查上年底稿和模板的列位置是否一致
"""

import os
import openpyxl

def check_alignment(template_path, prior_path, sheet_name, search_text="期末审定数"):
    """检查模板和上年底稿的列对齐情况"""
    results = {}

    # 检查模板
    try:
        wb_template = openpyxl.load_workbook(template_path, data_only=True)
        if sheet_name in wb_template.sheetnames:
            ws = wb_template[sheet_name]
            for row in range(1, min(80, ws.max_row + 1)):
                for col in range(1, 20):
                    val = ws.cell(row=row, column=col).value
                    if val and search_text in str(val):
                        results["template"] = {"row": row, "col": col}
                        break
                if "template" in results:
                    break
        wb_template.close()
    except Exception as e:
        results["template_error"] = str(e)

    # 检查上年底稿
    try:
        wb_prior = openpyxl.load_workbook(prior_path, data_only=True)
        if sheet_name in wb_prior.sheetnames:
            ws = wb_prior[sheet_name]
            for row in range(1, min(80, ws.max_row + 1)):
                for col in range(1, 20):
                    val = ws.cell(row=row, column=col).value
                    if val and search_text in str(val):
                        results["prior"] = {"row": row, "col": col}
                        break
                if "prior" in results:
                    break
        wb_prior.close()
    except Exception as e:
        results["prior_error"] = str(e)

    return results


# 测试所有科目
subjects = [
    ("C", "C.00 Lead", "C SWP 货币资金 202YMMDD XYZ公司.xlsx", "C SWP 货币资金 20251231 东风鸿远工程咨询有限公司.xlsx"),
    ("J1", "J.00  Lead Sheet", "J1 SWP 在建工程 202YMMDD XYZ公司.xlsx", "J1 SWP 在建工程 20251231 浙江力积-ok.xlsx"),
    ("K1", "K.00 Lead Sheet", "K1 SWP 固定资产 202YMMDD XYZ公司.xlsx", "输入：K1 固定资产 20251231六六六有限公司.xlsx"),
    ("L1", "L1.00 Lead sheet", "L1 SWP 无形资产 202YMMDD XYZ公司.xlsx", "L1 SWP 无形资产 20250831 东风鸿远工程咨询有限公司.xlsx"),
    ("L2", "L2.00 Lead", "L2 SWP 长期待摊费用 202YMMDD XYZ公司.xlsx", "L2 SWP 长期待摊费用 20251231 东风鸿远工程咨询有限公司.xlsx"),
    ("M", "M.00 Lead", "M SWP 应付票据 202YMMDD XYZ公司.xlsx", "M SWP 应付票据 20251231 DPCA.xlsx"),
    ("N", "N.00 Lead sheet", "N SWP  应付账款 202YMMDD XYZ公司.xlsx", "N SWP 应付账款 20251231 东风鸿远工程咨询有限公司.xlsx"),
    ("Q1", "Q1.00 Lead", "Q1 SWP 银行借款 202YMMDD XYZ公司.xlsx", "Q1 SWP 银行借款 20251231 东风鸿远工程咨询有限公司.xlsx"),
    ("Uexp", "Uexp_Lead", "U_exp SWP other 202YMMDD XYZ公司.xlsx", "U_exp SWP other 财务费用 20241231 YN.xlsx"),
]

base_dir = r"F:\AI专用\AI工作区\底稿roll forward"

for code, sheet_name, template_file, prior_file in subjects:
    template_path = os.path.join(base_dir, "templates", template_file)
    prior_path = os.path.join(base_dir, "上年底稿文件夹", prior_file)

    if not os.path.exists(prior_path):
        # 也尝试在根目录查找
        prior_path = os.path.join(base_dir, prior_file)

    if not os.path.exists(template_path):
        print(f"\n{code}: 模板文件不存在: {template_file}")
        continue

    if not os.path.exists(prior_path):
        print(f"\n{code}: 上年底稿不存在: {prior_file}")
        continue

    result = check_alignment(template_path, prior_path, sheet_name)

    print(f"\n{code} ({sheet_name}):")
    if "template" in result:
        t = result["template"]
        print(f"  模板: 行{t['row']}, 列{t['col']}")
    if "prior" in result:
        p = result["prior"]
        print(f"  上年: 行{p['row']}, 列{p['col']}")
    if "template" in result and "prior" in result:
        if result["template"]["col"] == result["prior"]["col"]:
            print(f"  [OK] 列号一致")
        else:
            print(f"  [WARN] 列号不一致！模板C{result['template']['col']} vs 上年C{result['prior']['col']}")
