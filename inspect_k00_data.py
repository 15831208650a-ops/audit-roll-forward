import openpyxl
import json

# 检查 K.00 Lead Sheet 数据区域（第30行以后）
prior_path = "输入：K1 固定资产 20251231六六六有限公司.xlsx"
wb = openpyxl.load_workbook(prior_path, data_only=False)
ws = wb["K.00 Lead Sheet"]

result = []
for row in range(30, min(70, ws.max_row + 1)):
    for col in range(1, min(15, ws.max_column + 1)):
        val = ws.cell(row=row, column=col).value
        if val:
            result.append({'row': row, 'col': col, 'value': str(val)[:100]})

with open('inspect_k00_data.json', 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

wb.close()
print("Done")
