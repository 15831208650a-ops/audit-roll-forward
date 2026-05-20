import openpyxl

# 验证K.01的公式
output_path = "output_v2/K1 SWP 固定资产 20261231 六六六有限公司.xlsx"
wb = openpyxl.load_workbook(output_path, data_only=False)
ws = wb["K.01 Agree SL to GL"]

print("=== K.01 第12行公式验证 ===")
print("列5(房屋建筑物账面数):", ws.cell(row=12, column=5).value)
print("列7(房屋建筑物审定数):", ws.cell(row=12, column=7).value)
print("列8(机器设备账面数):", ws.cell(row=12, column=8).value)
print("列10(机器设备审定数):", ws.cell(row=12, column=10).value)
print("列17(合计账面数):", ws.cell(row=12, column=17).value)
print("列19(合计审定数):", ws.cell(row=12, column=19).value)

# 验证第13行（购置）是否有数据
print("\n=== K.01 第13行（购置）验证 ===")
print("列5:", ws.cell(row=13, column=5).value)
print("列8:", ws.cell(row=13, column=8).value)

wb.close()
