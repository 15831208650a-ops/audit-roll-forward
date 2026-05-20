#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
找到U_exp财务费用子表的数据行起始位置
"""

import openpyxl

def find_data_rows(filepath, sheet_name):
    """查找财务费用子表的数据行"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[sheet_name]

    print(f"Sheet: {sheet_name}\n")

    # 读取更多行来找到数据
    for row in range(31, min(60, ws.max_row + 1)):
        has_data = False
        row_data = []
        for col in range(1, min(20, ws.max_column + 1)):
            val = ws.cell(row=row, column=col).value
            if val and str(val).strip() and str(val).strip() not in ['1', '/T1']:
                has_data = True
                row_data.append(f"C{col}={str(val)[:60]}")
        if has_data:
            print(f"R{row}: {', '.join(row_data)}")

    wb.close()


if __name__ == "__main__":
    filepath = r"F:\AI专用\AI工作区\底稿roll forward\templates\U_exp SWP other 202YMMDD XYZ公司.xlsx"
    find_data_rows(filepath, "Uexp财务费用BKD")
