#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
调试项目名称编码
"""

import openpyxl

new_path = r"F:\AI专用\AI工作区\底稿roll forward\templates\C SWP 货币资金 202YMMDD XYZ公司.xlsx"

wb_new = openpyxl.load_workbook(new_path, data_only=False)
ws_new = wb_new["C.00 Lead"]

# 检查R62C2
val = ws_new.cell(row=62, column=2).value
print(f"值: {repr(val)}")
print(f"类型: {type(val)}")
print(f"长度: {len(str(val)) if val else 0}")

if val:
    print(f"包含'项目': {'项目' in str(val)}")
    print(f"包含'名称': {'名称' in str(val)}")
    print(f"包含'项目名称': {'项目名称' in str(val)}")
    print(f"UTF-8 bytes: {str(val).encode('utf-8')}")

    # 也检查搜索条件
    search_term = "项目名称"
    print(f"搜索词: {repr(search_term)}")
    print(f"搜索词bytes: {search_term.encode('utf-8')}")

# 直接搜索
print("\n搜索所有C2单元格:")
for row in range(60, 70):
    v = ws_new.cell(row=row, column=2).value
    if v:
        match = "项目名称" in str(v)
        print(f"  R{row}: match={match}, val={repr(v)}")

wb_new.close()
