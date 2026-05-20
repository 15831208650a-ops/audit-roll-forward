#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
找到U_exp财务费用子表中"上年期末审定数"和"本期期初审定数"的具体列位置
"""

import openpyxl

def find_expense_columns(filepath, sheet_name):
    """查找财务费用子表的列位置"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[sheet_name]

    print(f"Sheet: {sheet_name}\n")

    # 查找R29-R30的表头
    print("查找表头（R29-R30）:")
    for row in [29, 30]:
        print(f"\nRow {row}:")
        for col in range(1, min(20, ws.max_column + 1)):
            val = ws.cell(row=row, column=col).value
            if val:
                print(f"  C{col}: {str(val)[:80]}")

    # 读取数据行（R31开始的几行）
    print("\n\n数据行预览（R31-R35）:")
    for row in range(31, min(36, ws.max_row + 1)):
        print(f"\nRow {row}:")
        for col in range(1, min(20, ws.max_column + 1)):
            val = ws.cell(row=row, column=col).value
            if val and str(val).strip():
                print(f"  C{col}: {str(val)[:80]}")

    wb.close()


if __name__ == "__main__":
    filepath = r"F:\AI专用\AI工作区\底稿roll forward\templates\U_exp SWP other 202YMMDD XYZ公司.xlsx"
    find_expense_columns(filepath, "Uexp财务费用BKD")
