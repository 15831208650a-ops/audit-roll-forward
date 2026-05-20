#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
调试For Disclosure处理
"""

import openpyxl

# 检查新文件
output_path = r"F:\AI专用\AI工作区\底稿roll forward\templates\C SWP 货币资金 202YMMDD XYZ公司.xlsx"
wb = openpyxl.load_workbook(output_path, data_only=False)
ws = wb["C.00 Lead"]

print("=== 新文件For Disclosure调试 ===")
for row in range(1, ws.max_row + 1):
    for col in range(1, 5):
        val = ws.cell(row=row, column=col).value
        if val and "Disclosure" in str(val):
            print(f"找到Disclosure在R{row}C{col}: {val}")
            # 显示附近行
            for r in range(row, min(row + 10, ws.max_row + 1)):
                row_data = []
                for c in range(1, 8):
                    v = ws.cell(row=r, column=c).value
                    if v is not None:
                        row_data.append(f"C{c}={v}")
                if row_data:
                    print(f"  R{r}: {', '.join(row_data)}")
            break

wb.close()

# 检查旧文件
prior_path = r"F:\AI专用\AI工作区\底稿roll forward\上年底稿文件夹\C SWP 货币资金 20251231 东风鸿远工程咨询有限公司.xlsx"
wb2 = openpyxl.load_workbook(prior_path, data_only=True)
ws2 = wb2["C.00 Lead"]

print("\n=== 旧文件For Disclosure调试 ===")
for row in range(1, ws2.max_row + 1):
    for col in range(1, 5):
        val = ws2.cell(row=row, column=col).value
        if val and "Disclosure" in str(val):
            print(f"找到Disclosure在R{row}C{col}: {val}")
            # 显示附近行
            for r in range(row, min(row + 10, ws2.max_row + 1)):
                row_data = []
                for c in range(1, 8):
                    v = ws2.cell(row=r, column=c).value
                    if v is not None:
                        row_data.append(f"C{c}={v}")
                if row_data:
                    print(f"  R{r}: {', '.join(row_data)}")
            break

wb2.close()
