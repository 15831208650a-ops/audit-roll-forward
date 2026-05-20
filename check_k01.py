import openpyxl

# 检查输出文件 K.01
output_path = "output/K1 SWP 固定资产 20261231 六六六有限公司.xlsx"
wb = openpyxl.load_workbook(output_path, data_only=True)
ws = wb["K.01 Agree SL to GL"]

print("=== K.01 第12行（年初余额）数据 ===")
for col in range(3, 22):
    val = ws.cell(row=12, column=col).value
    if val is not None:
        print(f"  列{col}: {val}")

print("\n=== K.01 第13行（购置）数据 ===")
for col in range(3, 22):
    val = ws.cell(row=13, column=col).value
    if val is not None:
        print(f"  列{col}: {val}")

wb.close()
