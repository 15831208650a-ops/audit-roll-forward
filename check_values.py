import openpyxl

# 检查上年底稿 I 列的值
prior_path = "输入：K1 固定资产 20251231六六六有限公司.xlsx"
wb_prior = openpyxl.load_workbook(prior_path, data_only=True)
ws_prior = wb_prior["K.00 Lead Sheet"]

print("=== 上年底稿 K.00 Lead Sheet I 列（期末审定数）值 ===")
for row in range(48, 60):
    val = ws_prior.cell(row=row, column=9).value
    print(f"  I{row}: {val}")

print("\n=== 上年底稿 K.00 Lead Sheet J 列（上期末审定数）值 ===")
for row in range(48, 60):
    val = ws_prior.cell(row=row, column=10).value
    print(f"  J{row}: {val}")

wb_prior.close()

# 检查输出文件 J 列的值
output_path = "output/K1 SWP 固定资产 20261231 六六六有限公司.xlsx"
wb_new = openpyxl.load_workbook(output_path, data_only=True)
ws_new = wb_new["K.00 Lead Sheet"]

print("\n=== 新底稿 K.00 Lead Sheet J 列（期初数）值 ===")
for row in range(48, 60):
    val = ws_new.cell(row=row, column=10).value
    print(f"  J{row}: {val}")

wb_new.close()
