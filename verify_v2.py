import openpyxl

# 验证新输出文件
output_path = "output_v2/K1 SWP 固定资产 20261231 六六六有限公司.xlsx"
print(f"=== 验证输出文件: {output_path} ===\n")

wb = openpyxl.load_workbook(output_path, data_only=False)

# 1. 验证 Lead Sheet
ws = wb["K.00 Lead Sheet"]
print("=== K.00 Lead Sheet 验证 ===")
print(f"  C2 (公司名称): {ws.cell(row=2, column=3).value}")
print(f"  C3 (日期): {ws.cell(row=3, column=3).value}")
print(f"  C5 (TE): {ws.cell(row=5, column=3).value}")
print(f"  C6 (SAD): {ws.cell(row=6, column=3).value}")

# 验证期初数（J列）
print("\n=== 期初数验证（J列，第48行表头后）===")
for row in range(49, 60):
    val = ws.cell(row=row, column=10).value  # J列
    if val is not None:
        print(f"  J{row}: {val}")

# 2. 验证 K.01
ws_k01 = wb["K.01 Agree SL to GL"]
print("\n=== K.01 第10行（固定资产类别表头） ===")
for col in range(1, 22):
    val = ws_k01.cell(row=10, column=col).value
    if val:
        print(f"  ({col}): {val}")

print("\n=== K.01 第12行（年初余额） ===")
for col in range(3, 22):
    val = ws_k01.cell(row=12, column=col).value
    if val is not None:
        print(f"  ({col}): {val}")

wb.close()
