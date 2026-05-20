#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查C货币资金输出文件的数据
"""

import openpyxl

def inspect_C_output(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet_name = "C.00 Lead"
    ws = wb[sheet_name]

    print(f"检查文件: {filepath}\n")

    # 1. 检查第3行（日期）
    print("=== 第3行（日期）===")
    for col in range(1, 6):
        val = ws.cell(row=3, column=col).value
        if val:
            print(f"  C{col}: {val}")

    # 2. 检查第5行（TE）
    print("\n=== 第5行（TE）===")
    for col in range(1, 6):
        val = ws.cell(row=5, column=col).value
        if val:
            print(f"  C{col}: {val}")

    # 3. 检查第6行（SAD）
    print("\n=== 第6行（SAD）===")
    for col in range(1, 6):
        val = ws.cell(row=6, column=col).value
        if val:
            print(f"  C{col}: {val}")

    # 4. 检查第14行（CRA）
    print("\n=== 第14行（CRA）===")
    for col in range(1, 6):
        val = ws.cell(row=14, column=col).value
        if val:
            print(f"  C{col}: {val}")

    # 5. 检查第37-38行（表头行）
    print("\n=== 第37行（表头行）===")
    for col in range(1, 16):
        val = ws.cell(row=37, column=col).value
        if val:
            print(f"  C{col}: {val}")

    print("\n=== 第38行（数据行）===")
    for col in range(1, 16):
        val = ws.cell(row=38, column=col).value
        if val is not None:
            print(f"  C{col}: {val}")

    # 6. 检查第60-62行（For Disclosure表2）
    print("\n=== 第60-62行（For Disclosure）===")
    for row in range(60, 63):
        print(f"Row {row}:")
        for col in range(1, 12):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                print(f"  C{col}: {val}")

    wb.close()

if __name__ == "__main__":
    inspect_C_output(r"F:\AI专用\AI工作区\底稿roll forward\output_final\C SWP 货币资金 20261231 东风鸿远工程咨询有限公司.xlsx")
