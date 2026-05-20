#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
调试max_row
"""

import openpyxl

new_path = r"F:\AI专用\AI工作区\底稿roll forward\templates\C SWP 货币资金 202YMMDD XYZ公司.xlsx"

wb_new = openpyxl.load_workbook(new_path, data_only=False)
ws_new = wb_new["C.00 Lead"]

print(f"max_row: {ws_new.max_row}")
print(f"max_column: {ws_new.max_column}")

for row in range(60, 70):
    val = ws_new.cell(row=row, column=2).value
    if val:
        print(f"R{row}C2: {val}")

wb_new.close()
