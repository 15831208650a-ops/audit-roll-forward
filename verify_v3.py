#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
验证v3版本的Roll Forward结果
"""

import os
import openpyxl


def verify_output(output_path):
    """验证输出文件"""
    wb = openpyxl.load_workbook(output_path, data_only=True)
    result = {"filename": os.path.basename(output_path), "sheets": {}}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if "Lead" in sheet_name:
            lead_info = {}
            # 查找客户名称
            for row in range(1, 10):
                for col in range(1, 6):
                    val = ws.cell(row=row, column=col).value
                    if val and "六六六" in str(val):
                        lead_info["company_name"] = str(val)
                        lead_info["company_location"] = f"R{row},C{col}"

            # 查找"期末审定数"所在行
            header_row = None
            for row in range(1, 80):
                for col in range(1, 15):
                    val = ws.cell(row=row, column=col).value
                    if val and "期末审定数" in str(val):
                        header_row = row
                        lead_info["header_row"] = row
                        break
                if header_row:
                    break

            # 读取期初数列数据
            if header_row:
                for col in range(1, 15):
                    val = ws.cell(row=header_row, column=col).value
                    if val and ("上期期末审定数" in str(val) or "期初数" in str(val)):
                        lead_info["opening_col"] = col
                        # 读取几行数据
                        data_preview = {}
                        for r in range(header_row + 1, min(header_row + 5, ws.max_row + 1)):
                            cell_val = ws.cell(row=r, column=col).value
                            if cell_val is not None:
                                data_preview[f"R{r}"] = str(cell_val)[:50]
                        lead_info["opening_data"] = data_preview
                        break

            result["sheets"][sheet_name] = lead_info

        elif "Agree" in sheet_name:
            k01_info = {}
            # 查找"年初余额"行
            for row in range(1, 20):
                for col in range(1, 5):
                    val = ws.cell(row=row, column=col).value
                    if val and "年初余额" in str(val):
                        k01_info["opening_balance_row"] = row
                        break
                if "opening_balance_row" in k01_info:
                    break

            # 读取表头行
            for row in range(8, 13):
                for col in range(1, 20):
                    val = ws.cell(row=row, column=col).value
                    if val and "固定资产类别" in str(val):
                        k01_info["header_row"] = row
                        break
                if "header_row" in k01_info:
                    break

            result["sheets"][sheet_name] = k01_info

    wb.close()
    return result


def main():
    output_dir = r"F:\AI专用\AI工作区\底稿roll forward\output_v3"
    if not os.path.exists(output_dir):
        print(f"输出目录不存在: {output_dir}")
        return

    files = [f for f in os.listdir(output_dir) if f.endswith('.xlsx')]
    for f in sorted(files):
        filepath = os.path.join(output_dir, f)
        result = verify_output(filepath)
        print(f"\n文件: {result['filename']}")
        for sheet, info in result['sheets'].items():
            print(f"  [{sheet}]: {info}")


if __name__ == "__main__":
    main()
