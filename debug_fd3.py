#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
调试For Disclosure处理3
"""

import openpyxl

prior_path = r"F:\AI专用\AI工作区\底稿roll forward\上年底稿文件夹\C SWP 货币资金 20251231 东风鸿远工程咨询有限公司.xlsx"
new_path = r"F:\AI专用\AI工作区\底稿roll forward\templates\C SWP 货币资金 202YMMDD XYZ公司.xlsx"

wb_prior = openpyxl.load_workbook(prior_path, data_only=True)
ws_prior = wb_prior["C.00 Lead"]

wb_new = openpyxl.load_workbook(new_path, data_only=False)
ws_new = wb_new["C.00 Lead"]

# 1. 旧文件R71行内容
print("=== 旧文件R71 ===")
for col in range(1, 8):
    val = ws_prior.cell(row=71, column=col).value
    print(f"  C{col}: {val} (type={type(val).__name__})")

# 2. 新文件R62行内容
print("\n=== 新文件R62 ===")
for col in range(1, 8):
    val = ws_new.cell(row=62, column=col).value
    print(f"  C{col}: {val} (type={type(val).__name__})")

# 3. 检查"项目名称"搜索
print("\n=== 搜索'项目名称' ===")
for row in range(1, 80):
    val = ws_new.cell(row=row, column=2).value
    if val and "项目名称" in str(val):
        print(f"  找到在R{row}: {val}")

wb_prior.close()
wb_new.close()
